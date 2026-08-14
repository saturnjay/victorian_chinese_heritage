#!/usr/bin/env python3
"""Build an Excel audit comparing Scattered Legacy with the VCH web map."""

from __future__ import annotations

import argparse
import io
import json
import re
import sys
import textwrap
import urllib.parse
import urllib.request
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps


SITE = "https://www.scatteredlegacy.org.au"
S3_PREFIX = "https://scatteredlegacy.s3.ap-southeast-2.amazonaws.com"
AUDIT_DATE = "2026-07-30"

VICTORIAN_REGIONS = {
    "Central Goldfields (Vic)",
    "Gippsland",
    "Melbourne",
    "Ovens Goldfields",
    "Victoria",
    "Western Districts (Vic)",
}
VICTORIAN_RIVERINA_PLACES = {"Chiltern", "Echuca", "Rutherglen", "Wahgunyah", "Wodonga"}

# Scattered Legacy records already represented by a discrete map record or a
# clearly equivalent map feature.
EXISTING_MATCHES = {
    "pig-oven-majorca": "VCH-043",
    "canton-lead-ararat": "VCH-061",
    "bi-jiang-grave": "VCH-037",
    "burning-tower-avoca": "VCH-021",
    "brickmaking-kiln": "VCH-005",
    "burning-tower-ballarat-new-cemetery": "VCH-025",
    "burning-tower-ballarat-old-cemetery": "VCH-026",
    "golden-point-joss-house-ballarat": "VCH-054",
    "ironbark-joss-house-ballarat": "VCH-056",
    "bendigo-joss-house-temple": "VCH-014",
    "burning-tower-bendigo": "VCH-007",
    "golden-dragon-museum-": "VCH-059",
    "loong-bendigo-dragon": "VCH-013",
    "burning-tower-castlemaine": "VCH-003",
    "castlemaine-joss-house": "VCH-055",
    "burning-tower-dunolly": "VCH-022",
    "harrietville-joss-house": "VCH-040",
    "burning-tower-maldon": "VCH-009",
    "burning-tower-maryborough": "VCH-028",
    "mount-misery-joss-house": "VCH-042",
    "pig-oven-mount-misery": "VCH-042",
    "chinese-section-vaughan-cemetery": "VCH-002 / VCH-031",
    "jericho-joss-house": "VCH-041",
    "pig-oven-jericho": "VCH-041",
    "burning-tower-melbourne-general-cemetery": "VCH-027",
    "chinese-mission-church-melbourne": "VCH-017",
    "kuomintang-kmt--nationalist-party-building-melbourne": "VCH-064",
    "museum-of-chinese-australian-history-melbourne": "VCH-058",
    "num-pon-soon-building": "VCH-016",
    "see-yup-temple--melbourne": "VCH-012",
    "beechworth-cemetery": "VCH-008",
    "burning-towers-beechworth": "VCH-008",
    "chinese-protectors-office-beechworth": "VCH-052",
    "memorial-beechworth-cemetery": "VCH-008",
    "bright-chinese-camp": "VCH-011",
    "bright-joss-house": "VCH-011",
    "buckland-cemetery": "VCH-034",
    "buckland-riot-memorial": "VCH-034",
    "joss-house-hill-burial-ground-cemetery": "VCH-033",
    "indigo-cornishtown-joss-house": "VCH-051",
    "yackandandah-cemetery-": "VCH-024",
    "burning-tower-echuca": "VCH-038",
    "burning-towers-carlyle-cemetery": "VCH-020",
    "memorial-chinese-burials-carlyle-cemetery": "VCH-020",
    "georgies-hut-buloke-shire": "VCH-019",
}

# High-value findings that appear absent as discrete map records.
STRONG_CANDIDATES = {
    "avoca-chinese-garden": (
        "High",
        "Add as a verified public interpretation/memorial place after confirming its current municipal status and coordinates.",
    ),
    "kong-meng-historic-reserve": (
        "High",
        "Add as an interpreted mining landscape associated with Lowe Kong Meng; confirm the reserve boundary and managing authority.",
    ),
    "thomas-coto-gravestone": (
        "High",
        "Add as a named grave at Sale Cemetery after confirming plot location, inscription, and cemetery access guidance.",
    ),
    "chinese-section-wangaratta-cemetery": (
        "High",
        "Add the Chinese and Strangers section after checking cemetery plans and present grave-marker condition.",
    ),
    "pioneer-cemetery-stanley-": (
        "High",
        "Add the marked pioneer-cemetery site after verifying the Chinese burial evidence and exact public access point.",
    ),
    "myrtleford-cemetery": (
        "High",
        "Add the Chinese section and community memorial after confirming its location within the cemetery.",
    ),
    "chiltern-cemetery": (
        "High",
        "Add the surviving Chinese graves after confirming the section, grave count, and cemetery authority.",
    ),
    "bowmans-forest-cemetery": (
        "High",
        "Add the Chinese and Strangers section; corroborate the seven burials and four reported surviving stones.",
    ),
    "alluvial-goldfields-stanley": (
        "Medium",
        "Investigate as a Chinese-worked mining landscape around Chinaman Road; define an evidence-based extent before mapping.",
    ),
    "harry-goons-hut-wodonga": (
        "High",
        "Add after confirming whether the photographed hut/camp survives, its exact location, ownership, and access sensitivity.",
    ),
    "creswick-chinese-camp": (
        "Medium",
        "Add as the former Chinese camp/Calembeen Park interpretation landscape if local heritage evidence confirms the mapped extent.",
    ),
}

# Selected records whose images are embedded in the workbook. This includes all
# strong candidates plus representative archival leads.
KEY_PHOTO_SLUGS = [
    "avoca-chinese-garden",
    "kong-meng-historic-reserve",
    "thomas-coto-gravestone",
    "chinese-section-wangaratta-cemetery",
    "pioneer-cemetery-stanley-",
    "myrtleford-cemetery",
    "chiltern-cemetery",
    "bowmans-forest-cemetery",
    "alluvial-goldfields-stanley",
    "harry-goons-hut-wodonga",
    "creswick-chinese-camp",
    "ararat-joss-house",
    "sale-joss-house",
    "little-bourke-street-joss-house",
    "omeo-joss-house",
    "stanley-joss-house",
]

QUESTIONABLE = {
    "baths-amherst": (
        "Do not add — disputed attribution",
        "Low",
        "The Scattered Legacy entry itself says there is little evidence for a Chinese bath house and suggests the explanation may be invented.",
    ),
    "silk-cloth-with-poem-by-liang-qichao": (
        "Exclude — non-Victorian/misfiled",
        "Low",
        "The linked place is Orange, NSW; the Victorian regional tag appears erroneous.",
    ),
    "foo-dog": (
        "Do not map — location unknown",
        "Low",
        "The pair is only said to be in Melbourne and its present whereabouts are unknown.",
    ),
    "cistern": (
        "Do not add — disputed attribution",
        "Low",
        "The entry cautions that ordinary cisterns are commonly misattributed as Chinese-built wells.",
    ),
}

DOCUMENTARY_JOSS = {
    "ararat-joss-house",
    "avoca-joss-house",
    "bairnsdale-joss-house",
    "beechworth-joss-house",
    "blackwood-joss-house",
    "buckland-joss-house",
    "chiltern-joss-house",
    "creswick-joss-house",
    "daylesford-joss-house",
    "germantown-joss-house",
    "granite-flat-joss-house",
    "growlers-creek-joss-house",
    "guildford-joss-house",
    "heathcote-joss-house",
    "hepburn-joss-house",
    "little-bourke-street-joss-house",
    "loo-toy-joss-house-north-bendigo",
    "majorca-joss-house",
    "maryborough-joss-house",
    "omeo-joss-house",
    "rutherglen-joss-house",
    "sale-joss-house",
    "stanley-joss-house",
    "vaughan-joss-house",
}


def plain_text(node) -> str:
    """Flatten Payload/Lexical rich text without retaining markup."""
    if isinstance(node, dict):
        pieces = []
        if isinstance(node.get("text"), str):
            pieces.append(node["text"])
        for key, value in node.items():
            if key != "text":
                pieces.append(plain_text(value))
        return " ".join(filter(None, pieces))
    if isinstance(node, list):
        return " ".join(filter(None, (plain_text(value) for value in node)))
    return ""


def clean(value, limit: int | None = None) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if limit and len(text) > limit:
        text = text[: limit - 1].rstrip() + "…"
    if text.startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text


def relation_title(value) -> str:
    if isinstance(value, dict):
        return clean(value.get("title") or value.get("name"))
    return clean(value)


def names(value) -> str:
    if not value:
        return ""
    values = value if isinstance(value, list) else [value]
    return "; ".join(
        clean(item.get("name") or item.get("title")) if isinstance(item, dict) else clean(item)
        for item in values
        if item
    )


def parse_coordinates(place: dict | None) -> tuple[float | None, float | None]:
    if not isinstance(place, dict):
        return None, None
    raw = place.get("location (longitude, latitude)")
    if not isinstance(raw, str) or "," not in raw:
        return None, None
    try:
        longitude, latitude = (float(part.strip()) for part in raw.split(",", 1))
        return latitude, longitude
    except ValueError:
        return None, None


def is_victorian(post: dict) -> bool:
    region = relation_title(post.get("region"))
    place = relation_title(post.get("place"))
    return region in VICTORIAN_REGIONS or (
        region == "Riverina (NSW/VIC)" and place in VICTORIAN_RIVERINA_PLACES
    )


def source_links(node) -> list[str]:
    links: list[str] = []
    if isinstance(node, dict):
        fields = node.get("fields")
        if isinstance(fields, dict) and isinstance(fields.get("url"), str):
            links.append(fields["url"])
        for value in node.values():
            links.extend(source_links(value))
    elif isinstance(node, list):
        for value in node:
            links.extend(source_links(value))
    return list(dict.fromkeys(link for link in links if link.startswith(("http://", "https://"))))


def media_items(node, role="Content image") -> list[dict]:
    found: list[dict] = []
    if isinstance(node, dict):
        if node.get("blockType") == "mediaBlock" and isinstance(node.get("media"), dict):
            item = dict(node["media"])
            item["_role"] = role
            found.append(item)
        for value in node.values():
            found.extend(media_items(value, role))
    elif isinstance(node, list):
        for value in node:
            found.extend(media_items(value, role))
    return found


def media_url(media: dict | None) -> str:
    if not isinstance(media, dict):
        return ""
    external = media.get("externalUrl")
    if isinstance(external, str) and external.startswith(("http://", "https://")):
        return external
    thumbnail = media.get("thumbnailURL")
    if isinstance(thumbnail, str) and thumbnail.startswith(("http://", "https://")):
        return urllib.parse.quote(thumbnail, safe=":/?=&%")
    url = media.get("url")
    if isinstance(url, str):
        return urllib.parse.urljoin(SITE, url)
    return ""


def media_caption(media: dict | None) -> str:
    if not isinstance(media, dict):
        return ""
    return clean(plain_text(media.get("caption")))


def media_resolution_score(media: dict) -> int:
    width = media.get("width")
    height = media.get("height")
    if isinstance(width, (int, float)) and isinstance(height, (int, float)):
        return int(width * height)
    url = media_url(media)
    width_hint = re.search(r"(\d{3,4})w\b", url)
    if width_hint:
        return int(width_hint.group(1)) ** 2
    dimensions = re.search(r"/(\d{2,4})/(\d{2,4})(?:-|_)", url)
    if dimensions:
        return int(dimensions.group(1)) * int(dimensions.group(2))
    return 0


def classify(post: dict) -> tuple[str, str, str, str]:
    slug = post["slug"]
    if slug in EXISTING_MATCHES:
        note = "Equivalent or broader coverage already exists on the map."
        if slug == "brickmaking-kiln":
            note += " Scattered Legacy labels the place as Ballarat, but its text describes the North Bendigo Fok Sing kiln (VCH-005)."
        return "Existing map coverage", "—", note, EXISTING_MATCHES[slug]
    if slug in STRONG_CANDIDATES:
        priority, recommendation = STRONG_CANDIDATES[slug]
        return "Missing candidate — investigate/add", priority, recommendation, ""
    if slug in QUESTIONABLE:
        assessment, priority, note = QUESTIONABLE[slug]
        return assessment, priority, note, ""
    if slug in DOCUMENTARY_JOSS:
        return (
            "Missing documentary lead — verify before mapping",
            "Medium" if slug in {
                "ararat-joss-house",
                "bairnsdale-joss-house",
                "creswick-joss-house",
                "little-bourke-street-joss-house",
                "omeo-joss-house",
                "stanley-joss-house",
            } else "Low",
            "The entry documents a former joss house, but a discrete map point should wait for a defensible location, relationship to any broader mapped site, and evidence of surviving archaeology or interpretation.",
            "",
        )
    nature = names(post.get("natureOfTreasure")).lower()
    if "pig oven (description only)" in nature:
        return (
            "Context/material — no verified surviving place",
            "Low",
            "Useful historical evidence, but the record is a description rather than a located surviving feature.",
            "",
        )
    return (
        "Context/material — enrich an existing or future record",
        "—",
        "Retain as documentary or collection context; it is not presently a well-defined separate Victorian map place.",
        "",
    )


def parse_map_records(index_path: Path) -> list[dict]:
    text = index_path.read_text(encoding="utf-8")
    match = re.search(r"const DATA = (\[.*?\]);\s*\n", text)
    if not match:
        raise RuntimeError("Could not locate the base DATA array in index.html")
    records = json.loads(match.group(1))
    block_start = text.find("    DATA.push(", match.end())
    block_end = text.find(");\n    const GROUP_COLORS", block_start)
    if block_start != -1 and block_end != -1:
        additions = json.loads(
            "["
            + text[block_start + len("    DATA.push("):block_end].replace(
                ");\n    DATA.push(", ","
            )
            + "]"
        )
        for addition in additions:
            if not any(record["ID"] == addition["ID"] for record in records):
                records.append(addition)
    return sorted(records, key=lambda record: record["ID"])


def build_rows(posts: list[dict], places: list[dict]) -> tuple[list[dict], list[dict]]:
    place_by_id = {place["id"]: place for place in places}
    rows = []
    photos = []
    for post in sorted((post for post in posts if is_victorian(post)), key=lambda p: (
        relation_title(p.get("region")), relation_title(p.get("place")), clean(p.get("title"))
    )):
        place_ref = post.get("place")
        place = place_by_id.get(place_ref.get("id"), place_ref) if isinstance(place_ref, dict) else {}
        latitude, longitude = parse_coordinates(place)
        assessment, priority, recommendation, map_match = classify(post)
        hero = post.get("heroImage") if isinstance(post.get("heroImage"), dict) else None
        all_media = []
        if hero:
            hero_copy = dict(hero)
            hero_copy["_role"] = "Hero image"
            all_media.append(hero_copy)
        all_media.extend(media_items(post.get("content")))
        deduped_media = []
        seen_media = set()
        for media in all_media:
            identity = media.get("id") or media_url(media)
            if identity and identity not in seen_media:
                seen_media.add(identity)
                deduped_media.append(media)
        key_media = max(deduped_media, key=media_resolution_score, default=hero or {})
        page_url = f"{SITE}/posts/{post['slug']}"
        links = source_links(post.get("content"))
        row = {
            "Assessment": assessment,
            "Priority": priority,
            "Map match ID": map_match,
            "Title": clean(post.get("title")),
            "Place": relation_title(post.get("place")),
            "Region": relation_title(post.get("region")),
            "Year / period": clean(post.get("year")),
            "Nature of treasure": names(post.get("natureOfTreasure")),
            "Keywords": names(post.get("keywords")),
            "Latitude": latitude,
            "Longitude": longitude,
            "Published": clean(post.get("publishedAt")),
            "Updated": clean(post.get("updatedAt")),
            "Scattered Legacy URL": page_url,
            "Summary / evidence": clean(plain_text(post.get("content")), 5000),
            "Recommendation / comparison note": recommendation,
            "Source links in entry": "\n".join(links),
            "Hero image URL": media_url(hero),
            "Hero image caption": media_caption(hero),
            "Hero image courtesy": clean((hero or {}).get("Image Courtesy of")),
            "Key image URL": media_url(key_media),
            "Key image caption": media_caption(key_media),
            "Key image courtesy": clean((key_media or {}).get("Image Courtesy of")),
            "Image count captured": len(deduped_media),
            "Scattered Legacy ID": clean(post.get("id")),
            "Slug": clean(post.get("slug")),
        }
        rows.append(row)
        for index, media in enumerate(deduped_media, 1):
            photos.append({
                "Treasure title": row["Title"],
                "Place": row["Place"],
                "Assessment": assessment,
                "Page URL": page_url,
                "Image number": index,
                "Role": media.get("_role", "Content image"),
                "Image name": clean(media.get("name") or media.get("filename")),
                "Caption": media_caption(media),
                "Courtesy / rights holder named": clean(media.get("Image Courtesy of")),
                "Image URL": media_url(media),
                "External source URL": clean(media.get("externalUrl")),
                "Mime type": clean(media.get("mimeType")),
                "Width": media.get("width"),
                "Height": media.get("height"),
                "Media ID": clean(media.get("id")),
                "Rights note": "Check with the named rights holder or collecting institution before reuse.",
                "Slug": row["Slug"],
            })
    return rows, photos


def download_key_image(row: dict, destination: Path) -> tuple[Path | None, str]:
    url = row.get("Key image URL", "")
    if not url:
        return None, "No hero image URL supplied"
    request = urllib.request.Request(url, headers={
        "User-Agent": "Mozilla/5.0 (research workbook; contact via project repository)",
        "Referer": row["Scattered Legacy URL"],
    })
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            payload = response.read()
        with Image.open(io.BytesIO(payload)) as source:
            image = ImageOps.exif_transpose(source).convert("RGB")
            image.thumbnail((720, 440), Image.Resampling.LANCZOS)
            canvas = Image.new("RGB", (720, 440), "white")
            left = (canvas.width - image.width) // 2
            top = (canvas.height - image.height) // 2
            canvas.paste(image, (left, top))
            canvas.save(destination, "JPEG", quality=88, optimize=True)
        return destination, "Downloaded and embedded for research reference"
    except Exception as error:  # Keep the workbook usable if a third-party host blocks downloading.
        return None, f"Download failed: {error}"


HEADER_FILL = PatternFill("solid", fgColor="8B1E2D")
SUBHEADER_FILL = PatternFill("solid", fgColor="E9D7D9")
ACCENT_FILL = PatternFill("solid", fgColor="F7F1E8")
HIGH_FILL = PatternFill("solid", fgColor="E2F0D9")
MEDIUM_FILL = PatternFill("solid", fgColor="FFF2CC")
LOW_FILL = PatternFill("solid", fgColor="FCE4D6")
GREY_FILL = PatternFill("solid", fgColor="E7E6E6")
WHITE_FONT = Font(color="FFFFFF", bold=True)
THIN_GREY = Side(style="thin", color="D9D9D9")


def style_table(ws, widths: dict[str, float] | None = None, freeze="A2"):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN_GREY)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN_GREY)
    if widths:
        for column, width in widths.items():
            ws.column_dimensions[column].width = width


def add_rows_sheet(wb: Workbook, title: str, rows: list[dict], columns: list[str], widths: list[float]):
    ws = wb.create_sheet(title)
    ws.append(columns)
    for row in rows:
        ws.append([row.get(column, "") for column in columns])
    style_table(ws, {get_column_letter(i + 1): width for i, width in enumerate(widths)})
    url_columns = {i + 1 for i, column in enumerate(columns) if "URL" in column or column == "Source links in entry"}
    for row in range(2, ws.max_row + 1):
        assessment = str(ws.cell(row, 1).value or "")
        priority = str(ws.cell(row, 2).value or "") if ws.max_column >= 2 else ""
        if assessment.startswith("Missing candidate"):
            ws.cell(row, 1).fill = HIGH_FILL
        elif assessment.startswith("Missing documentary"):
            ws.cell(row, 1).fill = MEDIUM_FILL
        elif assessment.startswith("Existing"):
            ws.cell(row, 1).fill = GREY_FILL
        elif assessment.startswith(("Do not", "Exclude")):
            ws.cell(row, 1).fill = LOW_FILL
        if priority == "High":
            ws.cell(row, 2).fill = HIGH_FILL
        elif priority == "Medium":
            ws.cell(row, 2).fill = MEDIUM_FILL
        elif priority == "Low":
            ws.cell(row, 2).fill = LOW_FILL
        for column in url_columns:
            cell = ws.cell(row, column)
            if isinstance(cell.value, str) and cell.value.startswith("http") and "\n" not in cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
    return ws


def build_workbook(rows: list[dict], photos: list[dict], map_records: list[dict], output: Path, image_dir: Path):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:F1")
    summary["A1"] = "Scattered Legacy → Victorian Chinese Heritage Map Audit"
    summary["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    summary["A1"].fill = HEADER_FILL
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 38
    summary.merge_cells("A3:F3")
    summary["A3"] = (
        "A research comparison of Scattered Legacy’s public Victorian catalogue against the "
        f"{len(map_records)} findings currently loaded by the interactive map."
    )
    summary["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    summary.row_dimensions[3].height = 42

    counts = Counter(row["Assessment"] for row in rows)
    metrics = [
        ("Audit date", AUDIT_DATE),
        ("Scattered Legacy treasures reviewed", len(rows)),
        ("Current map findings", len(map_records)),
        ("Strong missing candidates", counts["Missing candidate — investigate/add"]),
        ("Documentary leads", counts["Missing documentary lead — verify before mapping"]),
        ("Existing map coverage", counts["Existing map coverage"]),
        ("Context/material records", sum(v for k, v in counts.items() if k.startswith("Context/material"))),
        ("Questionable/excluded records", sum(v for k, v in counts.items() if k.startswith(("Do not", "Exclude")))),
        ("Image metadata rows", len(photos)),
        ("Embedded key-photo records", len(KEY_PHOTO_SLUGS)),
    ]
    summary.append([])
    summary.append(["Metric", "Value"])
    for metric, value in metrics:
        summary.append([metric, value])
    for cell in summary[5]:
        cell.fill = SUBHEADER_FILL
        cell.font = Font(bold=True)

    start = summary.max_row + 3
    summary.cell(start, 1, "Key conclusion")
    summary.cell(start, 1).font = Font(bold=True, color="FFFFFF")
    summary.cell(start, 1).fill = HEADER_FILL
    summary.merge_cells(start_row=start, start_column=1, end_row=start, end_column=6)
    conclusion = (
        "Scattered Legacy identifies multiple plausible omissions. The clearest additions are public or extant "
        "places: Avoca Chinese Garden, Kong Meng Historic Reserve, Thomas Coto’s grave, Chinese cemetery sections "
        "at Wangaratta, Myrtleford, Chiltern and Bowmans Forest, Stanley’s pioneer cemetery and mining landscape, "
        "Harry Goon’s Hut/camp in Wodonga, and the former Creswick Chinese camp landscape. Numerous former joss "
        "houses are retained as documentary leads because their exact locations and surviving evidence require "
        "corroboration before map publication."
    )
    summary.merge_cells(start_row=start + 1, start_column=1, end_row=start + 3, end_column=6)
    summary.cell(start + 1, 1, conclusion)
    summary.cell(start + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    summary.cell(start + 1, 1).fill = ACCENT_FILL

    rights_row = start + 5
    summary.cell(rights_row, 1, "Image-rights warning")
    summary.cell(rights_row, 1).font = Font(bold=True, color="FFFFFF")
    summary.cell(rights_row, 1).fill = HEADER_FILL
    summary.merge_cells(start_row=rights_row, start_column=1, end_row=rights_row, end_column=6)
    rights = (
        "Scattered Legacy states that third-party images may be used with permission, under an open licence, "
        "believed to be public domain, or linked from an original source. Inclusion does not mean an image is "
        "freely reusable. Check each named rights holder or institution before publication. Embedded thumbnails "
        "in this workbook are research-reference copies and retain page, image URL, caption and courtesy fields."
    )
    summary.merge_cells(start_row=rights_row + 1, start_column=1, end_row=rights_row + 3, end_column=6)
    summary.cell(rights_row + 1, 1, rights)
    summary.cell(rights_row + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    summary.cell(rights_row + 1, 1).fill = LOW_FILL
    summary["A" + str(rights_row + 5)] = "Primary catalogue"
    summary["B" + str(rights_row + 5)] = SITE
    summary["B" + str(rights_row + 5)].hyperlink = SITE
    summary["B" + str(rights_row + 5)].style = "Hyperlink"
    summary["A" + str(rights_row + 6)] = "Rights/disclaimer"
    summary["B" + str(rights_row + 6)] = f"{SITE}/disclaimers"
    summary["B" + str(rights_row + 6)].hyperlink = f"{SITE}/disclaimers"
    summary["B" + str(rights_row + 6)].style = "Hyperlink"
    summary["A" + str(rights_row + 7)] = "Project description"
    summary["B" + str(rights_row + 7)] = f"{SITE}/about"
    summary["B" + str(rights_row + 7)].hyperlink = f"{SITE}/about"
    summary["B" + str(rights_row + 7)].style = "Hyperlink"
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 28
    for column in "CDEF":
        summary.column_dimensions[column].width = 18

    candidates = [
        row for row in rows
        if row["Assessment"].startswith(("Missing candidate", "Missing documentary"))
    ]
    candidate_columns = [
        "Assessment", "Priority", "Title", "Place", "Region", "Year / period",
        "Latitude", "Longitude", "Nature of treasure", "Summary / evidence",
        "Recommendation / comparison note", "Scattered Legacy URL", "Hero image URL",
        "Hero image courtesy", "Published", "Updated", "Slug",
    ]
    add_rows_sheet(
        wb, "Missing candidates", candidates, candidate_columns,
        [34, 10, 34, 18, 22, 15, 12, 12, 25, 64, 60, 42, 42, 28, 20, 20, 28],
    )

    inventory_columns = [
        "Assessment", "Priority", "Map match ID", "Title", "Place", "Region",
        "Year / period", "Nature of treasure", "Keywords", "Latitude", "Longitude",
        "Published", "Updated", "Scattered Legacy URL", "Summary / evidence",
        "Recommendation / comparison note", "Source links in entry", "Hero image URL",
        "Hero image caption", "Hero image courtesy", "Image count captured",
        "Scattered Legacy ID", "Slug",
    ]
    add_rows_sheet(
        wb, "Victorian treasure audit", rows, inventory_columns,
        [34, 10, 15, 34, 18, 22, 15, 25, 34, 12, 12, 20, 20, 42, 64, 58, 48, 42, 32, 28, 12, 28, 30],
    )

    photo_columns = [
        "Treasure title", "Place", "Assessment", "Page URL", "Image number", "Role",
        "Image name", "Caption", "Courtesy / rights holder named", "Image URL",
        "External source URL", "Mime type", "Width", "Height", "Media ID", "Rights note", "Slug",
    ]
    photo_ws = add_rows_sheet(
        wb, "Photo metadata", photos, photo_columns,
        [34, 18, 34, 42, 10, 16, 32, 40, 32, 46, 46, 14, 10, 10, 28, 48, 30],
    )
    for row_index in range(2, photo_ws.max_row + 1):
        for column in (4, 10, 11):
            cell = photo_ws.cell(row_index, column)
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"

    key_ws = wb.create_sheet("Key photos")
    key_ws.sheet_view.showGridLines = False
    key_ws.append([
        "Preview", "Finding / material", "Assessment", "Place", "Caption",
        "Courtesy / rights holder named", "Scattered Legacy page", "Image URL",
        "Local research copy", "Download status", "Rights note",
    ])
    style_table(key_ws, {
        "A": 54, "B": 34, "C": 34, "D": 18, "E": 38, "F": 30,
        "G": 42, "H": 46, "I": 42, "J": 36, "K": 48,
    })
    row_by_slug = {row["Slug"]: row for row in rows}
    image_dir.mkdir(parents=True, exist_ok=True)
    for slug in KEY_PHOTO_SLUGS:
        row = row_by_slug.get(slug)
        if not row:
            continue
        destination = image_dir / f"{slug}.jpg"
        image_path, status = download_key_image(row, destination)
        key_ws.append([
            "", row["Title"], row["Assessment"], row["Place"], row["Key image caption"],
            row["Key image courtesy"], row["Scattered Legacy URL"], row["Key image URL"],
            str(image_path.relative_to(output.parent)) if image_path else "", status,
            "Research-reference copy only. Confirm reuse permission with the named rights holder/institution.",
        ])
        current_row = key_ws.max_row
        key_ws.row_dimensions[current_row].height = 220
        if image_path:
            image = XLImage(str(image_path))
            image.width = 360
            image.height = 220
            key_ws.add_image(image, f"A{current_row}")
        for column in (7, 8):
            cell = key_ws.cell(current_row, column)
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"

    map_ws = wb.create_sheet("Current map index")
    map_ws.append(["Map ID", "Finding name"])
    for record in map_records:
        map_ws.append([record.get("ID"), record.get("Name")])
    style_table(map_ws, {"A": 14, "B": 62})

    method = wb.create_sheet("Methodology")
    method.sheet_view.showGridLines = False
    methodology_rows = [
        ("Scope", "All Scattered Legacy treasures assigned to Central Goldfields (Vic), Gippsland, Melbourne, Ovens Goldfields, Victoria, Western Districts (Vic), plus the Victorian places Chiltern, Echuca, Rutherglen, Wahgunyah and Wodonga within Riverina (NSW/VIC)."),
        ("Catalogue snapshot", f"Public Scattered Legacy API examined on {AUDIT_DATE}: 576 treasures and 192 places in the full catalogue; 112 treasures passed the Victorian scope rule."),
        ("Map baseline", f"The local interactive map currently loads {len(map_records)} findings: the base DATA array plus records appended at runtime."),
        ("Comparison method", "Titles, places, treasure types, summaries, coordinates and related media were compared with map names, locality, scope and known broader-site coverage. Matching is interpretive, not solely string-based."),
        ("Strong candidate rule", "A seemingly absent, distinct Victorian place with extant fabric, a public memorial/interpretive landscape, a named grave/cemetery section, or otherwise strong location-specific evidence."),
        ("Documentary lead rule", "A former site—principally joss houses—with historical evidence but insufficient exact location, surviving fabric, statutory status or present access information for immediate publication."),
        ("Context/material rule", "Portable objects, documents, newspapers, collection portals, generic historical images, or evidence that enriches a mapped place without constituting a separate map location."),
        ("Caution", "Scattered Legacy describes itself as an ongoing research tool and notes that many entries are partial or incomplete. Candidate findings require corroboration with statutory, municipal, cemetery, archaeological or custodian sources."),
        ("Image rights", "The site’s disclaimer states that third-party media is not automatically free to reuse. The workbook retains the stated courtesy and URLs; embedded images are research-reference copies."),
        ("Known metadata issue", "The Scattered Legacy 'Brickmaking kiln' treasure is assigned to Ballarat, but its text describes the North Bendigo Fok Sing kiln already mapped as VCH-005."),
        ("Known regional issue", "The Liang Qichao silk-cloth treasure is assigned to Orange, NSW despite a Central Goldfields (Vic) region tag; it is excluded as a Victorian map candidate."),
        ("Next research step", "Prioritise the High candidates, obtain authoritative coordinates and current-condition evidence, then decide whether documentary joss-house leads should be separate points or annotations within broader settlement records."),
    ]
    method.append(["Topic", "Method / note"])
    for topic, note in methodology_rows:
        method.append([topic, note])
    style_table(method, {"A": 28, "B": 120})
    for row in range(2, method.max_row + 1):
        method.row_dimensions[row].height = 58

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def validate_workbook(path: Path, expected_map_rows: int):
    workbook = load_workbook(path, read_only=False, data_only=False)
    required = {
        "Summary", "Missing candidates", "Victorian treasure audit",
        "Photo metadata", "Key photos", "Current map index", "Methodology",
    }
    missing = required.difference(workbook.sheetnames)
    if missing:
        raise RuntimeError(f"Workbook validation failed; missing sheets: {sorted(missing)}")
    if workbook["Victorian treasure audit"].max_row != 113:
        raise RuntimeError("Workbook validation failed; Victorian audit should contain 112 data rows")
    if workbook["Current map index"].max_row != expected_map_rows + 1:
        raise RuntimeError(
            f"Workbook validation failed; current map index should contain {expected_map_rows} data rows"
        )
    workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--posts", type=Path, required=True)
    parser.add_argument("--places", type=Path, required=True)
    parser.add_argument("--index", type=Path, default=Path("index.html"))
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    posts_payload = json.loads(args.posts.read_text(encoding="utf-8"))
    places_payload = json.loads(args.places.read_text(encoding="utf-8"))
    posts = posts_payload["docs"]
    places = places_payload["docs"]
    rows, photos = build_rows(posts, places)
    map_records = parse_map_records(args.index)
    image_dir = args.output.parent / "scattered_legacy_key_photos"
    build_workbook(rows, photos, map_records, args.output, image_dir)
    validate_workbook(args.output, len(map_records))

    counts = Counter(row["Assessment"] for row in rows)
    print(f"Workbook: {args.output}")
    print(f"Victorian treasures: {len(rows)}")
    print(f"Strong missing candidates: {counts['Missing candidate — investigate/add']}")
    print(f"Documentary leads: {counts['Missing documentary lead — verify before mapping']}")
    print(f"Existing map coverage: {counts['Existing map coverage']}")
    print(f"Photo metadata rows: {len(photos)}")
    print(f"Embedded key images: {len(load_workbook(args.output)['Key photos']._images)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
